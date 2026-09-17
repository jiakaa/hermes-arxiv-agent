#!/usr/bin/env python3
"""一次性:给 papers_record.xlsx 的 Papers 表末尾追加 rating / keywords 两列(空值)。幂等。"""
import sys
from pathlib import Path
from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent.parent
EXCEL = BASE / "papers_record.xlsx"
NEW_COLS = ["rating", "keywords"]


def main() -> int:
    wb = load_workbook(EXCEL)
    ws = wb["Papers"]
    header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    header = [str(h) if h is not None else "" for h in header]
    added = []
    for name in NEW_COLS:
        if name not in header:
            ws.cell(row=1, column=len(header) + 1, value=name)
            header.append(name)
            added.append(name)
    if added:
        wb.save(EXCEL)
        print(f"[OK] Added columns: {added}; header now {len(header)} cols: {header}")
    else:
        print("[OK] Columns already present, nothing to do.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
