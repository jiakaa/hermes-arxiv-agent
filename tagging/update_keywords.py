#!/usr/bin/env python3
"""规则化打标签:vocabulary.json 的别名命中 title+abstract+summary_cn。

--missing 只处理 keywords 为空的论文(日常增量);--all 全量重打;--dry-run 只看不改。
写回 Excel keywords 列(分号分隔,每篇最多 8 个,按词表顺序截断)。幂等。
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent.parent
EXCEL = BASE / "papers_record.xlsx"
VOCAB = Path(__file__).resolve().parent / "vocabulary.json"
MAX_TAGS = 8
ASCII_ONLY = re.compile(r"^[a-z0-9 \-/'._]+$")


def load_vocab() -> list[tuple[str, str, list[str], list[str], list[str]]]:
    """-> [(group, tag, ascii_single_tokens, ascii_phrases, cjk_aliases)]"""
    d = json.loads(VOCAB.read_text(encoding="utf-8"))
    out = []
    for g in d["groups"]:
        for t in g["tags"]:
            singles, phrases, cjk = [], [], []
            for a in t.get("aliases", []):
                if not ASCII_ONLY.match(a.lower()):
                    cjk.append(a)
                elif " " in a.strip() or "-" in a.strip():
                    phrases.append(a.lower())
                else:
                    singles.append(a.lower())
            out.append((g["name"], t["name"], singles, phrases, cjk))
    return out


def tag_paper(title: str, abstract: str, summary_cn: str, vocab) -> list[str]:
    text = f"{title}\n{abstract}\n{summary_cn}"
    low = text.lower()
    hits = []
    for _group, name, singles, phrases, cjk in vocab:
        # 单词别名要求词边界,避免 "CARD" 命中 "cards" 这类假阳;短语/中文按子串
        if (any(re.search(r"\b" + re.escape(a) + r"\b", low) for a in singles)
                or any(a in low for a in phrases)
                or any(a in text for a in cjk)):
            hits.append(name)
        if len(hits) >= MAX_TAGS:
            break
    return hits


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--missing", action="store_true", help="只处理 keywords 为空的论文")
    ap.add_argument("--all", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    if not (args.missing or args.all):
        ap.error("need --missing or --all")

    vocab = load_vocab()
    wb = load_workbook(EXCEL)
    ws = wb["Papers"]
    header = [str(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i + 1 for i, h in enumerate(header)}
    for need in ("arxiv_id", "title", "abstract", "summary_cn", "keywords"):
        assert need in idx, f"missing column {need}"

    dist, changed, total = Counter(), 0, 0
    for r in range(2, ws.max_row + 1):
        def get(name):
            v = ws.cell(row=r, column=idx[name]).value
            return "" if v is None else str(v)

        if not get("arxiv_id"):
            continue
        if args.missing and get("keywords"):
            continue
        tags = tag_paper(get("title"), get("abstract"), get("summary_cn"), vocab)
        total += 1
        joined = ";".join(tags)
        for t in tags:
            dist[t] += 1
        if get("keywords") != joined:
            changed += 1
            if not args.dry_run:
                ws.cell(row=r, column=idx["keywords"], value=joined)

    print(f"[INFO] processed={total} changed={changed} dry_run={args.dry_run}")
    print(f"[INFO] top tags: {dist.most_common(25)}")
    if not args.dry_run and changed:
        wb.save(EXCEL)
        print(f"[OK] keywords written, {changed} rows updated")
    return 0


if __name__ == "__main__":
    sys.exit(main())
