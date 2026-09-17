"""monitor.py Excel 写入的表头驱动行为测试(旧工作簿自动补列、值对齐、upsert 不重复行)。

运行:python3 -m unittest tests.test_excel_writes -v
"""
import sys
import unittest
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))

import monitor  # noqa: E402

OLD_HEADER = [
    "arxiv_id", "title", "authors", "affiliations", "published_date", "categories",
    "abstract", "summary_cn", "pdf_filename", "crawled_date", "notes",
]
NEW_HEADER = OLD_HEADER + ["rating", "keywords"]

PAPER = {
    "arxiv_id": "2609.99999",
    "title": "T",
    "authors": "A",
    "published_date": "2026-09-01",
    "categories": "cs.CV",
    "summary": "abs-text",
    "pdf_filename": "2609.99999.pdf",
    "affiliations": "Somewhere",
    "summary_cn": "中文总结",
    "rating": "S",
    "keywords": "3DGS;世界模型",
}


def make_wb(header):
    wb = openpyxl.Workbook()
    ws = wb[wb.sheetnames[0]]
    ws.title = "Papers"
    ws.append(header)
    return wb


class TestAppendToExcel(unittest.TestCase):
    def test_old_workbook_gains_new_columns_with_aligned_values(self):
        wb = make_wb(OLD_HEADER)
        monitor.append_to_excel(wb, PAPER)
        ws = wb["Papers"]
        header = [c.value for c in ws[1]]
        self.assertEqual(header, NEW_HEADER, "旧表头应自动补出 rating/keywords")
        row = [c.value for c in ws[2]]
        self.assertEqual(row[header.index("rating")], "S")
        self.assertEqual(row[header.index("keywords")], "3DGS;世界模型")
        self.assertEqual(row[header.index("abstract")], "abs-text")
        self.assertEqual(row[header.index("authors")], "A")
        self.assertEqual(len(row), len(header), "行长度必须与表头一致")

    def test_new_workbook_does_not_duplicate_columns(self):
        wb = make_wb(NEW_HEADER)
        monitor.append_to_excel(wb, PAPER)
        header = [c.value for c in wb["Papers"][1]]
        self.assertEqual(header, NEW_HEADER)
        self.assertEqual(len(header), len(set(header)))

    def test_upsert_updates_rating_keywords_without_appending_row(self):
        wb = make_wb(NEW_HEADER)
        monitor.append_to_excel(wb, {**PAPER, "rating": "", "keywords": ""})
        ws = wb["Papers"]
        header_index, row_index = monitor.build_excel_row_index(ws)
        monitor.upsert_to_excel(
            ws, header_index, row_index, {**PAPER, "rating": "A", "keywords": "NeRF"}
        )
        self.assertEqual(ws.max_row, 2, "upsert 不得新增重复行")
        row = [c.value for c in ws[2]]
        self.assertEqual(row[header_index["rating"] - 1], "A")
        self.assertEqual(row[header_index["keywords"] - 1], "NeRF")

    def test_upsert_does_not_clear_existing_values_with_empty_input(self):
        wb = make_wb(NEW_HEADER)
        monitor.append_to_excel(wb, PAPER)
        ws = wb["Papers"]
        header_index, row_index = monitor.build_excel_row_index(ws)
        monitor.upsert_to_excel(
            ws, header_index, row_index,
            {**PAPER, "rating": "", "keywords": "", "affiliations": "", "summary_cn": ""},
        )
        row = [c.value for c in ws[2]]
        self.assertEqual(row[header_index["rating"] - 1], "S", "空值不应覆盖已有评级")
        self.assertEqual(row[header_index["keywords"] - 1], "3DGS;世界模型")
        self.assertEqual(row[header_index["affiliations"] - 1], "Somewhere")


if __name__ == "__main__":
    unittest.main(verbosity=2)
